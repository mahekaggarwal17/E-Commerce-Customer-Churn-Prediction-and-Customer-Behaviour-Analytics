import os
import json
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def create_report():
    # Load model metadata for exact ground-truth values
    with open('models/model_metadata.json', 'r') as f:
        meta = json.load(f)
    
    results = meta['results']
    lr_res = results['Logistic Regression']
    rf_res = results['Random Forest Classifier']

    doc = docx.Document()

    # Page Margins: 1 inch on all sides
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Styles setup
    style_normal = doc.styles['Normal']
    font_normal = style_normal.font
    font_normal.name = 'Calibri'
    font_normal.size = Pt(11)
    font_normal.color.rgb = RGBColor(40, 40, 40)

    # ==========================================
    # COVER PAGE
    # ==========================================
    p_title_space = doc.add_paragraph()
    p_title_space.paragraph_format.space_before = Pt(40)

    p_badge = doc.add_paragraph()
    p_badge.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_badge = p_badge.add_run("ACADEMIC INTERNSHIP PROJECT REPORT")
    r_badge.bold = True
    r_badge.font.size = Pt(13)
    r_badge.font.color.rgb = RGBColor(43, 92, 143)

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(20)
    p_title.paragraph_format.space_after = Pt(20)
    r_title = p_title.add_run("E-Commerce Customer Churn Prediction and Customer Behaviour Analytics")
    r_title.bold = True
    r_title.font.size = Pt(24)
    r_title.font.color.rgb = RGBColor(30, 58, 138)

    p_subtitle = doc.add_paragraph()
    p_subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sub = p_subtitle.add_run("A Machine Learning & Exploratory Data Analytics Implementation for Proactive Customer Retention")
    r_sub.italic = True
    r_sub.font.size = Pt(13)
    r_sub.font.color.rgb = RGBColor(100, 100, 100)

    p_div = doc.add_paragraph()
    p_div.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_div.paragraph_format.space_before = Pt(30)
    p_div.paragraph_format.space_after = Pt(40)
    r_div = p_div.add_run("—" * 35)
    r_div.font.color.rgb = RGBColor(200, 200, 200)

    # Submission Meta Table
    meta_table = doc.add_table(rows=6, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_table.autofit = False

    meta_data = [
        ("Candidate Name:", "[Your Name / StudentName]"),
        ("Program Name:", "IBM SkillsBuild Data Analytics with AI Academic Internship"),
        ("Conducting Organization:", "BharatCares"),
        ("Institutional Association:", "All India Council for Technical Education (AICTE)"),
        ("Domain Track:", "Data Analytics, Machine Learning & Business Intelligence"),
        ("Date of Submission:", "September 2026")
    ]

    for idx, (label, val) in enumerate(meta_data):
        row = meta_table.rows[idx]
        cell_lbl, cell_val = row.cells[0], row.cells[1]
        cell_lbl.width = Inches(2.2)
        cell_val.width = Inches(4.2)
        
        p_l = cell_lbl.paragraphs[0]
        r_l = p_l.add_run(label)
        r_l.bold = True
        r_l.font.color.rgb = RGBColor(30, 58, 138)
        
        p_v = cell_val.paragraphs[0]
        r_v = p_v.add_run(val)
        if "[" in val:
            r_v.italic = True
            r_v.font.color.rgb = RGBColor(180, 50, 50)
        else:
            r_v.font.color.rgb = RGBColor(50, 50, 50)
            
        set_cell_margins(cell_lbl, 80, 80, 100, 100)
        set_cell_margins(cell_val, 80, 80, 100, 100)

    doc.add_page_break()

    # ==========================================
    # HELPER FUNCTIONS FOR SECTIONS
    # ==========================================
    def add_sec_heading(title, level=1):
        h = doc.add_heading(title, level=level)
        h.paragraph_format.space_before = Pt(16)
        h.paragraph_format.space_after = Pt(8)
        run = h.runs[0]
        if level == 1:
            run.font.size = Pt(16)
            run.font.color.rgb = RGBColor(30, 58, 138)
            run.bold = True
        elif level == 2:
            run.font.size = Pt(13)
            run.font.color.rgb = RGBColor(43, 92, 143)
            run.bold = True
        return h

    def add_para(text, bold_prefix=None, space_after=6):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            r_pre = p.add_run(bold_prefix)
            r_pre.bold = True
            r_pre.font.color.rgb = RGBColor(30, 58, 138)
        p.add_run(text)
        return p

    def add_bullet(text, bold_prefix=None):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            r_pre = p.add_run(bold_prefix)
            r_pre.bold = True
        p.add_run(text)
        return p

    def embed_figure(image_path, caption_text, width=Inches(5.5)):
        if os.path.exists(image_path):
            p_img = doc.add_paragraph()
            p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_img.paragraph_format.space_before = Pt(10)
            p_img.paragraph_format.space_after = Pt(4)
            p_img.add_run().add_picture(image_path, width=width)
            
            p_cap = doc.add_paragraph()
            p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_cap.paragraph_format.space_after = Pt(12)
            r_cap = p_cap.add_run(caption_text)
            r_cap.font.size = Pt(9.5)
            r_cap.italic = True
            r_cap.font.color.rgb = RGBColor(100, 100, 100)

    # ==========================================
    # 1. ABSTRACT
    # ==========================================
    add_sec_heading("1. Executive Abstract")
    add_para(
        "In modern e-commerce enterprises, customer attrition represents a substantial impediment to sustained growth and lifetime value realization. "
        "Acquiring replacement consumers incurs five to twenty-five times greater capital expenditure than preserving established buyers. "
        "This project implements an end-to-end data analytics, behavioral segmentation, and machine learning framework for proactive customer churn prevention "
        "as part of the IBM SkillsBuild Data Analytics with AI Academic Internship Program conducted by BharatCares in association with AICTE."
    )
    add_para(
        "Utilizing a verified, publicly accessible e-commerce dataset comprising 5,630 customer transaction profiles across 20 distinct demographic, operational, "
        "and behavioral attributes, this investigation establishes a systematic analytical pipeline. The study employs Exploratory Data Analysis (EDA) to expose "
        "the relationship between attrition and customer friction indicators, implements an industry-standard RFM (Recency, Frequency, Monetary) segmentation model, "
        "and compares classification algorithms. The baseline Logistic Regression model achieves an accuracy of 79.22% with 84.74% recall, whereas the ensemble "
        "Random Forest Classifier demonstrates superior discriminative performance, achieving 96.98% accuracy, 98.95% recall, an F1-score of 0.9171, and an "
        "ROC-AUC of 0.9971 on stratified testing data. Empirical feature importance reveals that customer relationship duration (Tenure), customer service complaints, "
        "monetary cashback tiers, and warehouse fulfillment distances are the primary determinants of churn. To facilitate operational deployment, an interactive "
        "Streamlit decision support dashboard was developed, featuring automated KPI tracking, behavioral cohort drill-downs, and a live churn probability simulator."
    )

    # ==========================================
    # 2. INTRODUCTION
    # ==========================================
    add_sec_heading("2. Introduction")
    add_para(
        "The digital transformation of retail commerce has democratized product discovery and purchase convenience, but it has simultaneously lowered customer "
        "switching barriers. Consumers routinely transition between competing digital platforms in response to minor delivery delays, unaddressed customer service "
        "disputes, or promotional incentives. Consequently, contemporary e-commerce operators must pivot from purely transactional acquisition models to data-driven "
        "retention engineering."
    )
    add_para(
        "Customer churn analytics leverages historical purchasing trajectories, service interactions, and demographic indicators to forecast which consumers exhibit "
        "a high likelihood of dormancy or platform abandonment. By proactively diagnosing churn propensity prior to total disengagement, customer success teams "
        "can orchestrate surgical, personalized retention interventions—such as tailored loyalty rewards, priority dispute resolution, or targeted onboarding drip "
        "campaigns—thereby safeguarding customer lifetime value (CLV) and maximizing return on customer acquisition costs (CAC)."
    )

    # ==========================================
    # 3. PROBLEM STATEMENT
    # ==========================================
    add_sec_heading("3. Problem Statement")
    add_para(
        "Despite accumulating voluminous event streams, modern e-commerce platforms struggle with three fundamental operational challenges in retention management:"
    )
    add_bullet(" Absence of automated, real-time warning mechanisms to flag disengaging customers before attrition becomes irreversible.", "1. Detection Blindspots:")
    add_bullet(" Inability to isolate and quantify the specific behavioral triggers (e.g., product return friction, delayed delivery, unanswered grievances) causing attrition.", "2. Causal Ambiguity:")
    add_bullet(" Reliance on generic, blanket discount promotions that erode operational margins and fail to engage high-value dormant segments.", "3. Inefficient Intervention:")
    add_para(
        "Hence, there is a clear academic and industrial requirement for an interpretable, scalable predictive model paired with an intuitive analytical dashboard "
        "to empower commercial decision-makers to retain valuable customer relationships."
    )

    # ==========================================
    # 4. OBJECTIVES
    # ==========================================
    add_sec_heading("4. Project Objectives")
    add_para("The core objectives executed throughout this academic internship project include:")
    add_bullet("Ingest the real-world e-commerce customer dataset, inspect data hygiene, remediate missing records, harmonize categorical aliases, and screen for anomalies.", "1. Preprocessing & Quality Control:")
    add_bullet("Execute rigorous univariate, bivariate, and multivariate visual analytics to characterize customer distributions and assess retention determinants.", "2. Exploratory Data Analytics:")
    add_bullet("Construct an RFM (Recency, Frequency, Monetary) behavioral scoring model to categorize the customer base into actionable behavioral cohorts.", "3. Behavioral Customer Segmentation:")
    add_bullet("Architect an end-to-end scikit-learn machine learning pipeline with stratified train-test partitioning and class-weight balancing.", "4. Predictive Machine Learning:")
    add_bullet("Benchmark Logistic Regression against a tuned Random Forest Classifier using multi-dimensional metrics (Accuracy, Precision, Recall, F1, ROC-AUC).", "5. Rigorous Model Evaluation:")
    add_bullet("Extract Gini-impurity feature importances and logistic odds ratios to provide transparent interpretability for commercial stakeholders.", "6. Model Interpretability:")
    add_bullet("Build and deploy a responsive Streamlit web application providing live KPI visualization and single-customer churn risk inference.", "7. Interactive Dashboard Deployment:")

    # ==========================================
    # 5. DATASET DESCRIPTION
    # ==========================================
    add_sec_heading("5. Dataset Description")
    add_para(
        "The project is grounded upon the legitimate, publicly accessible 'E-Commerce Customer Churn and Retention' dataset published on Kaggle. "
        "The dataset represents genuine retail platform interactions with 5,630 unique customer instances and 20 attributes."
    )
    add_bullet("Kaggle Public Repository (Ankit Verma)", "Source Repository: ")
    add_bullet("https://www.kaggle.com/datasets/ankitverma2010/ecommerce-customer-churn-analysis-and-prediction", "Dataset URL: ")
    add_bullet("CC0: Public Domain (Educational and Commercial Exploration Permitted)", "Licensing: ")
    add_bullet("5,630 Records | 20 Attributes (14 Numerical, 5 Categorical, 1 Target)", "Dimensions: ")
    add_bullet("Churn (Binary: 0 = Retained, 1 = Churned)", "Target Attribute: ")

    doc.add_paragraph().paragraph_format.space_before = Pt(4)
    # Table of Features
    schema_table = doc.add_table(rows=1, cols=3)
    schema_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = schema_table.rows[0].cells
    hdr[0].text, hdr[1].text, hdr[2].text = "Variable Name", "Data Type", "Operational Definition"
    for c in hdr:
        set_cell_background(c, "1E3A8A")
        c.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
        c.paragraphs[0].runs[0].bold = True
        set_cell_margins(c, 100, 100, 120, 120)

    schema_data = [
        ("CustomerID", "Integer", "Unique customer identification key (50001 to 55630)"),
        ("Churn", "Integer (0/1)", "Target variable indicating whether the customer lapsed"),
        ("Tenure", "Float", "Duration in months the customer has engaged with the company"),
        ("PreferredLoginDevice", "Categorical", "Primary device used to authenticate (Mobile Phone, Computer)"),
        ("CityTier", "Integer (1-3)", "Geographic development tier of customer municipality"),
        ("WarehouseToHome", "Float", "Physical transit distance (km) between fulfillment hub and customer"),
        ("PreferredPaymentMode", "Categorical", "Payment method (Credit Card, Debit Card, UPI, COD, E-wallet)"),
        ("Gender", "Categorical", "Gender of the consumer (Male / Female)"),
        ("HourSpendOnApp", "Float", "Weekly hours logged across platform apps and browser"),
        ("NumberOfDeviceRegistered", "Integer", "Number of distinct devices paired with the account"),
        ("PreferedOrderCat", "Categorical", "Top purchase category (Laptop, Mobile, Fashion, Grocery, Others)"),
        ("SatisfactionScore", "Integer (1-5)", "Customer-submitted post-purchase service rating"),
        ("MaritalStatus", "Categorical", "Marital status (Single, Married, Divorced)"),
        ("NumberOfAddress", "Integer", "Total physical shipping destinations registered on account"),
        ("Complain", "Binary (0/1)", "Indicator of formal complaint filed within the preceding 30 days"),
        ("OrderAmountHikeFromlastYear", "Float", "Percentage growth in annual order value over previous period"),
        ("CouponUsed", "Float", "Count of promotional voucher redemption events in last month"),
        ("OrderCount", "Float", "Total count of successful purchase orders placed in last month"),
        ("DaySinceLastOrder", "Float", "Recency metric indicating days elapsed since last purchase"),
        ("CashbackAmount", "Float", "Average promotional monetary cashback earned in last month ($)")
    ]

    for v_name, v_type, v_desc in schema_data:
        r = schema_table.add_row().cells
        r[0].width, r[1].width, r[2].width = Inches(1.8), Inches(1.2), Inches(3.4)
        r[0].paragraphs[0].add_run(v_name).bold = True
        r[1].paragraphs[0].add_run(v_type)
        r[2].paragraphs[0].add_run(v_desc)
        for c in r:
            set_cell_margins(c, 70, 70, 100, 100)

    # ==========================================
    # 6. TECHNOLOGIES USED
    # ==========================================
    add_sec_heading("6. Technologies & Software Stack")
    add_para("The analytical pipeline and software artifacts were implemented strictly using production-grade open-source tools:")
    add_bullet("Python 3.11.16: High-performance core programming runtime environment.", "Programming Runtime: ")
    add_bullet("Pandas 3.0.6 & NumPy 2.4.6: High-throughput tabular data structures, array operations, and statistical transformations.", "Data Manipulation: ")
    add_bullet("Matplotlib 3.11.2 & Seaborn 0.13.2: Publication-quality statistical plotting and visual distributions.", "Visual Analytics: ")
    add_bullet("Scikit-Learn 1.9.1: Machine learning algorithms, imputation pipelines, preprocessing ColumnTransformers, and diagnostic metrics.", "Machine Learning: ")
    add_bullet("Plotly 7.1.0 & Streamlit 1.64.0: Modern interactive browser application engine for dynamic dashboards and live inference forms.", "Web Dashboard: ")
    add_bullet("Python-Docx 1.2.0: Programmatic synthesis of standardized Microsoft Word academic project reports.", "Documentation Engine: ")
    add_bullet("Jupyter Notebook & NBFormat: Reproducible, interactive computational notebook development.", "Notebook Environment: ")

    # ==========================================
    # 7. METHODOLOGY
    # ==========================================
    add_sec_heading("7. Research Methodology")
    add_para(
        "This project adheres to the standardized Cross-Industry Standard Process for Data Mining (CRISP-DM) framework, "
        "structured into an eight-stage engineering workflow:"
    )
    add_bullet("Dataset Acquisition & Schema Validation: Ingestion of source records and baseline data integrity inspection.", "Stage 1 — ")
    add_bullet("Data Preprocessing: Handling missing values, standardizing categorical variations, and validating unique identifiers.", "Stage 2 — ")
    add_bullet("Exploratory Data Analysis: Univariate profiling, bivariate hypothesis testing, and correlation matrix estimation.", "Stage 3 — ")
    add_bullet("Behavioral RFM Segmentation: Derivation of Recency, Frequency, and Monetary scores to generate five customer cohorts.", "Stage 4 — ")
    add_bullet("Feature Engineering & Leakage Isolation: Construction of ColumnTransformer pipelines with median imputation and scaling.", "Stage 5 — ")
    add_bullet("Supervised Model Benchmarking: Stratified 80/20 train-test partitioning; training Logistic Regression and Random Forest models.", "Stage 6 — ")
    add_bullet("Comprehensive Diagnostic Evaluation: Quantifying Accuracy, Precision, Recall, F1-score, Confusion Matrices, and ROC-AUC curves.", "Stage 7 — ")
    add_bullet("Interactive Dashboard & Operational Deployment: Productionizing the optimal pipeline within an interactive Streamlit UI.", "Stage 8 — ")

    # ==========================================
    # 8. DATA PREPROCESSING
    # ==========================================
    add_sec_heading("8. Data Preprocessing & Cleansing")
    add_para(
        "To ensure analytical validity and safeguard the machine learning algorithms against biased estimation, "
        "data cleaning was conducted systematically:"
    )
    add_bullet(
        "Examination of the dataset revealed exactly 5,630 unique `CustomerID` instances with zero exact duplicate records across the observation base.",
        "1. Deduplication: "
    )
    add_bullet(
        "Inspection detected duplicate semantic representations in key categorical fields: in `PreferredLoginDevice`, 'Phone' was unified with 'Mobile Phone'; "
        "in `PreferredPaymentMode`, 'CC' was mapped to 'Credit Card' and 'COD' to 'Cash on Delivery'; and in `PreferedOrderCat`, 'Mobile' was harmonized to 'Mobile Phone'.",
        "2. Categorical Harmonization: "
    )
    add_bullet(
        "Seven numerical attributes contained missing observations: `Tenure` (264 nulls, 4.69%), `WarehouseToHome` (251 nulls, 4.46%), "
        "`HourSpendOnApp` (255 nulls, 4.53%), `OrderAmountHikeFromlastYear` (265 nulls, 4.71%), `CouponUsed` (256 nulls, 4.55%), "
        "`OrderCount` (258 nulls, 4.58%), and `DaySinceLastOrder` (307 nulls, 5.45%). Rather than deleting incomplete instances (which would purge over 15% "
        "of the training sample), robust median imputation was deployed within the scikit-learn preprocessing pipeline.",
        "3. Missing Value Remediation: "
    )
    add_bullet(
        "Numerical variables were standardized via `StandardScaler` to achieve zero mean and unit variance for linear classifiers, "
        "while categorical attributes were one-hot encoded (`OneHotEncoder`) with unknown category handling.",
        "4. Feature Normalization & Encoding: "
    )

    # ==========================================
    # 9. EXPLORATORY DATA ANALYSIS
    # ==========================================
    add_sec_heading("9. Exploratory Data Analysis & Empirical Insights")
    add_para(
        "Exploratory visual analytics uncovered prominent behavioral divergences between retained and churned customer populations:"
    )

    embed_figure("reports/figures/fig1_churn_distribution.png", "Figure 1: Class Distribution of Customer Churn (Target Variable)")

    add_para(
        "Figure 1 illustrates the distribution of the target variable. Of the 5,630 recorded customers, 4,682 (83.16%) remained retained with the platform, "
        "while 948 (16.84%) churned. This 5:1 class imbalance necessitates cost-sensitive machine learning with balanced class weighting to ensure rare "
        "churn events are identified with high sensitivity."
    )

    embed_figure("reports/figures/fig2_tenure_vs_churn.png", "Figure 2: Kernel Density Estimation (KDE) of Customer Tenure by Churn Status")

    add_para(
        "Figure 2 illustrates the distribution of customer tenure. Customers who churned exhibit a starkly contracted tenure distribution, with a mean of "
        "only 3.38 months (median 1.0 month), compared to retained consumers who average 11.50 months (median 10.0 months). This highlights an extreme "
        "vulnerability window during the initial 90 to 120 days post-acquisition."
    )

    embed_figure("reports/figures/fig3_complain_vs_churn.png", "Figure 3: Churn Proportion Stratified by Customer Complaint Status")

    add_para(
        "As depicted in Figure 3, customer service complaints function as a potent churn accelerator. Customers who logged a complaint in the previous month "
        "exhibit an empirical churn rate of 31.67%, compared to only 10.93% for customers without recorded complaints—a nearly three-fold increase."
    )

    embed_figure("reports/figures/fig4_category_performance.png", "Figure 4: Churn Rate (%) across Preferred Order Categories")

    add_para(
        "Analysis across product categories (Figure 4) indicates substantial divergence in customer stickiness. Shoppers primarily purchasing 'Mobile Phone' "
        "exhibit the highest churn rate at 27.40%, whereas consumers purchasing 'Grocery' items churn at an extraordinarily low rate of 4.88%, demonstrating "
        "the powerful retention stickiness of routine consumable retail."
    )

    embed_figure("reports/figures/fig6_correlation_heatmap.png", "Figure 5: Pearson Correlation Matrix of Numerical Features with Customer Churn")

    add_para(
        "Figure 5 presents the correlation structure among numerical attributes. `Tenure` displays the strongest inverse correlation with churn (-0.35), "
        "followed by `CashbackAmount` (-0.15) and `DaySinceLastOrder` (-0.16). Conversely, `Complain` (+0.25) and `CityTier` (+0.08) exhibit positive correlations with churn."
    )

    # ==========================================
    # 10. CUSTOMER SEGMENTATION (RFM)
    # ==========================================
    add_sec_heading("10. Customer Segmentation via RFM Analysis")
    add_para(
        "To empower commercial marketing teams with granular behavioral cohorts, an RFM model was established based on quintile binning:"
    )
    add_bullet("Derived from `DaySinceLastOrder`. Customers ordering within recent days receive higher scores (5 to 1).", "Recency (R): ")
    add_bullet("Derived from `OrderCount` in the preceding period. Higher order volume yields higher scores (1 to 5).", "Frequency (F): ")
    add_bullet("Derived from `CashbackAmount` (proportional cashback spending reward proxy). Higher tiers receive higher scores (1 to 5).", "Monetary (M): ")

    add_para("The cumulative score categorizes customers into five distinct operational segments:")

    embed_figure("reports/figures/fig5_rfm_segments.png", "Figure 6: Customer Churn Rate across RFM Behavioral Segments")

    # Table of RFM Segments
    rfm_table = doc.add_table(rows=1, cols=6)
    rfm_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    rfm_hdr = rfm_table.rows[0].cells
    rfm_hdr[0].text, rfm_hdr[1].text, rfm_hdr[2].text, rfm_hdr[3].text, rfm_hdr[4].text, rfm_hdr[5].text = (
        "Customer Segment", "Customer Count", "Share (%)", "Avg Orders", "Avg Cashback ($)", "Empirical Churn Rate"
    )
    for c in rfm_hdr:
        set_cell_background(c, "1E3A8A")
        c.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
        c.paragraphs[0].runs[0].bold = True
        set_cell_margins(c, 80, 80, 80, 80)

    rfm_rows = [
        ("Loyal Customers", "2,065", "36.68%", "3.5", "$185.12", "11.33%"),
        ("Potential Loyalists", "1,429", "25.38%", "2.4", "$167.90", "24.35%"),
        ("Hibernating", "1,075", "19.10%", "2.0", "$158.45", "19.72%"),
        ("At Risk", "644", "11.44%", "3.4", "$188.30", "13.98%"),
        ("Champions", "417", "7.41%", "4.8", "$212.40", "15.35%")
    ]

    for seg, cnt, shr, ords, cb, ch_rt in rfm_rows:
        row_c = rfm_table.add_row().cells
        row_c[0].paragraphs[0].add_run(seg).bold = True
        row_c[1].paragraphs[0].add_run(cnt)
        row_c[2].paragraphs[0].add_run(shr)
        row_c[3].paragraphs[0].add_run(ords)
        row_c[4].paragraphs[0].add_run(cb)
        row_c[5].paragraphs[0].add_run(ch_rt).bold = True
        for c in row_c:
            set_cell_margins(c, 60, 60, 80, 80)

    add_para(
        "Crucially, the segmentation demonstrates that 'Potential Loyalists' exhibit the highest churn vulnerability (24.35%), "
        "indicating that while these consumers made recent purchases, their platform loyalty is fragile and requires active reinforcement."
    )

    # ==========================================
    # 11. MACHINE LEARNING WORKFLOW
    # ==========================================
    add_sec_heading("11. Machine Learning Modeling & Pipeline Architecture")
    add_para(
        "To establish a robust predictive engine, the feature matrix was decoupled from the target variable (`Churn`), "
        "and metadata fields (`CustomerID`, derived RFM score indices) were isolated to eliminate risk of data leakage. "
        "The remaining 18 raw features (13 numerical, 5 categorical) were partitioned into an 80% training sample (4,504 records) "
        "and a 20% holdout test sample (1,126 records) utilizing stratified sampling to guarantee identical class distribution."
    )
    add_para(
        "Two contrasting learning algorithms were benchmarked:"
    )
    add_bullet(
        "Logistic Regression (Balanced): Serves as a transparent linear baseline. Regularized with L2 penalty, solved via L-BFGS, "
        "and parameterized with `class_weight='balanced'` to offset the 5:1 class disparity.",
        "1. Baseline Classifier: "
    )
    add_bullet(
        "Random Forest Classifier (Balanced): An ensemble of 150 de-correlated decision trees constrained to a maximum depth of 12. "
        "It constructs non-linear decision boundaries through bagging and random feature subspace sampling.",
        "2. Advanced Ensemble: "
    )

    # ==========================================
    # 12. RESULTS & MODEL EVALUATION
    # ==========================================
    add_sec_heading("12. Model Evaluation & Benchmark Results")
    add_para(
        "Model evaluation was performed on the independent 1,126-customer test partition. In customer churn analytics, precision and recall "
        "carry asymmetric financial costs: missing a churner (False Negative) forfeits total future customer revenue, whereas contacting a healthy "
        "customer (False Positive) merely incurs a negligible promotional communication cost."
    )

    # Comparison Table
    bench_table = doc.add_table(rows=1, cols=6)
    bench_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    b_hdr = bench_table.rows[0].cells
    b_hdr[0].text, b_hdr[1].text, b_hdr[2].text, b_hdr[3].text, b_hdr[4].text, b_hdr[5].text = (
        "Classification Algorithm", "Accuracy", "Precision (Churn)", "Recall (Churn)", "F1-Score (Churn)", "ROC-AUC"
    )
    for c in b_hdr:
        set_cell_background(c, "1E3A8A")
        c.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)
        c.paragraphs[0].runs[0].bold = True
        set_cell_margins(c, 80, 80, 80, 80)

    perf_data = [
        ("Logistic Regression (Balanced)", f"{lr_res['accuracy']*100:.2f}%", f"{lr_res['precision']*100:.2f}%", f"{lr_res['recall']*100:.2f}%", f"{lr_res['f1_score']:.4f}", f"{lr_res['roc_auc']:.4f}"),
        ("Random Forest Classifier (Balanced)", f"{rf_res['accuracy']*100:.2f}%", f"{rf_res['precision']*100:.2f}%", f"{rf_res['recall']*100:.2f}%", f"{rf_res['f1_score']:.4f}", f"{rf_res['roc_auc']:.4f}")
    ]

    for model_n, acc_v, prec_v, rec_v, f1_v, auc_v in perf_data:
        r_cells = bench_table.add_row().cells
        r_cells[0].paragraphs[0].add_run(model_n).bold = True
        r_cells[1].paragraphs[0].add_run(acc_v)
        r_cells[2].paragraphs[0].add_run(prec_v)
        r_cells[3].paragraphs[0].add_run(rec_v).bold = True
        r_cells[4].paragraphs[0].add_run(f1_v)
        r_cells[5].paragraphs[0].add_run(auc_v).bold = True
        for c in r_cells:
            set_cell_margins(c, 70, 70, 80, 80)

    embed_figure("reports/figures/fig7_confusion_matrices.png", "Figure 7: Confusion Matrices for Logistic Regression and Random Forest Classifiers")

    add_para(
        "As revealed in the confusion matrices (Figure 7), out of 190 actual churners in the holdout evaluation sample:"
    )
    add_bullet("Correctly identified 161 churners (Recall = 84.74%), with 29 False Negatives and 205 False Positives.", "Logistic Regression: ")
    add_bullet("Correctly captured 188 churners (Recall = 98.95%), achieving only 2 False Negatives and 32 False Positives.", "Random Forest Classifier: ")

    embed_figure("reports/figures/fig8_roc_curves.png", "Figure 8: Receiver Operating Characteristic (ROC) Comparison Curves")

    add_para(
        "Figure 8 displays the ROC curves. The Random Forest Classifier achieves an extraordinary area under the curve of 0.9971, "
        "demonstrating near-perfect discriminative power across all decision thresholds compared to the linear baseline (0.8851)."
    )

    # ==========================================
    # 13. FEATURE IMPORTANCE & INTERPRETABILITY
    # ==========================================
    add_sec_heading("13. Feature Importance & Model Interpretability")
    add_para(
        "To ensure business stakeholders understand the rationale governing automated churn predictions, "
        "we extracted the mean decrease in Gini impurity across the Random Forest ensemble."
    )

    embed_figure("reports/figures/fig9_feature_importances.png", "Figure 9: Top 15 Feature Importances extracted from Random Forest")

    add_para("The top five predictive determinants identified by the Random Forest model are:")
    add_bullet("Accounts for 27.17% of total predictive weight. Tenure is by far the single most dominant barrier against attrition.", "1. Tenure (0.2717): ")
    add_bullet("Accounts for 8.91% of predictive weight. Reflects monetary spending tier; customers receiving lower promotional cashback exhibit higher churn risk.", "2. Cashback Amount (0.0891): ")
    add_bullet("Accounts for 7.03% of predictive weight. The primary operational catalyst causing abrupt customer defection.", "3. Complaint Raised (0.0703): ")
    add_bullet("Accounts for 5.78% of predictive weight. Greater shipping distances correlate with delivery lag and customer dissatisfaction.", "4. Warehouse-to-Home Distance (0.0578): ")
    add_bullet("Accounts for 5.48% of predictive weight. Extended inactivity intervals signal impending customer churn.", "5. Days Since Last Order (0.0548): ")

    # ==========================================
    # 14. INTERACTIVE STREAMLIT APPLICATION
    # ==========================================
    add_sec_heading("14. Interactive Web Application Architecture")
    add_para(
        "To operationalize the findings, a responsive web dashboard was created using Streamlit (`app.py`). "
        "The application provides real-time analytical capabilities structured into five dedicated modules:"
    )
    add_bullet("Displays headline metrics (Total Customers, Churn Rate, Active Customers, Average Tenure, Average Cashback) and high-level distribution charts.", "1. Executive Overview: ")
    add_bullet("Enables drill-downs into RFM cohort profiles, preferred order categories, and payment channel distributions.", "2. Customer Behaviour & RFM: ")
    add_bullet("Interactive visual analysis exploring tenure drop-offs, complaint impacts, city tier dynamics, and transit distance friction.", "3. Churn Driver Deep Dive: ")
    add_bullet("Displays comparative model benchmarks, confusion matrices, ROC curves, and dynamic horizontal feature importance charts.", "4. AI Model Evaluation: ")
    add_bullet("Interactive input form allowing customer relationship managers to simulate customer attributes, generate live churn probabilities via the trained pipeline, display risk gauge meters, and receive automated retention strategy recommendations.", "5. Live Churn Risk Predictor: ")

    # ==========================================
    # 15. KEY FINDINGS
    # ==========================================
    add_sec_heading("15. Key Analytical Findings")
    add_bullet("Tenure is the critical determinant of loyalty: customer attrition is intensely concentrated in the first 0 to 4 months (>40% churn). Once tenure surpasses 10 months, retention stabilizes above 90%.", "1. Onboarding Fragility:")
    add_bullet("Grievances act as an immediate churn trigger: customers logging complaints in the last 30 days churn at 31.67%, compared to only 10.93% for non-complainants.", "2. The Complaint Hazard:")
    add_bullet("Product categories dictate retention dynamics: 'Mobile Phone' purchasers exhibit elevated churn (27.40%), whereas 'Grocery' consumers churn at a modest 4.88%.", "3. Category Disparity:")
    add_bullet("Geography and logistics impact churn: Tier 3 cities face a 21.37% churn rate versus 14.51% in Tier 1, strongly correlated with warehouse distance.", "4. Fulfillment Friction:")
    add_bullet("Random Forest achieves production-grade predictive excellence: with 96.98% accuracy and 98.95% recall, the model captured 188 out of 190 at-risk customers.", "5. ML Efficacy:")

    # ==========================================
    # 16. BUSINESS RECOMMENDATIONS
    # ==========================================
    add_sec_heading("16. Actionable Business Recommendations")
    add_bullet("Deploy a dedicated 90-day welcoming journey incorporating structured milestone rewards, usage tutorials, and proactive check-ins to steer new buyers past the high-risk 4-month mark.", "1. First-90-Days Concierge Onboarding:")
    add_bullet("Establish an automated CRM escalation protocol that flags unresolved customer complaints within 2 hours, paired with instant compensatory store credit or express shipping tokens.", "2. Rapid Complaint Remediation Protocol:")
    add_bullet("Direct targeted engagement campaigns toward 'Potential Loyalists' (24.35% churn) with category-tailored bundles to transform single-purchase shoppers into habitual buyers.", "3. Tailored RFM Retention Journeys:")
    add_bullet("Mitigate delivery friction for remote consumers by partnering with localized third-party logistics (3PL) partners and establishing regional micro-fulfillment hubs.", "4. Regional Logistics Modernization:")
    add_bullet("Integrate the trained Random Forest pipeline into the commercial CRM to run automated weekly risk scoring, routing all accounts with predicted churn probability >60% directly to retention specialists.", "5. Automated Churn Scoring in CRM:")

    # ==========================================
    # 17. LIMITATIONS
    # ==========================================
    add_sec_heading("17. Project Limitations")
    add_bullet("The dataset reflects a static historical snapshot without multi-quarter time-series granularity, precluding dynamic survival analysis.", "1. Temporal Snapshot:")
    add_bullet("Product-level profit margins and customer lifetime value were approximated through cashback rather than net transaction ledger receipts.", "2. Margin Visibility:")
    add_bullet("Customer feedback was captured via a single 1-to-5 numeric score and complaint flag; qualitative text feedback was unavailable for NLP sentiment analysis.", "3. Qualitative Data:")

    # ==========================================
    # 18. FUTURE SCOPE
    # ==========================================
    add_sec_heading("18. Future Scope")
    add_bullet("Formulate time-to-event survival models (e.g., Cox Proportional Hazards) to forecast the precise expected lifespan and time until churn for every consumer.", "1. Survival Analysis:")
    add_bullet("Incorporate Natural Language Processing (NLP) models on customer support transcripts and social media feedback to detect customer dissatisfaction sentiment prior to formal complaint logging.", "2. Multimodal Sentiment Analysis:")
    add_bullet("Deploy the predictive pipeline within cloud infrastructure (e.g., Google Cloud Vertex AI / AWS SageMaker) with automated model drift monitoring and retraining pipelines.", "3. Cloud MLOps Architecture:")

    # ==========================================
    # 19. CONCLUSION
    # ==========================================
    add_sec_heading("19. Conclusion")
    add_para(
        "This project successfully developed and validated an end-to-end data analytics and machine learning solution for e-commerce customer churn prediction. "
        "Through data preprocessing, rigorous exploratory analysis, RFM segmentation, and algorithm benchmarking, the project demonstrated that customer tenure, "
        "service complaints, fulfillment distance, and spending rewards represent the foundational pillars of customer retention."
    )
    add_para(
        "The engineered Random Forest Classifier delivers enterprise-grade predictive efficacy (96.98% accuracy, 98.95% recall, 0.9971 ROC-AUC), providing commercial "
        "decision-makers with a dependable early warning mechanism. Delivered alongside an interactive Streamlit application and comprehensive academic documentation, "
        "this project satisfies all requirements for the IBM SkillsBuild Data Analytics with AI Academic Internship Program conducted by BharatCares in association with AICTE."
    )

    # ==========================================
    # 20. REFERENCES
    # ==========================================
    add_sec_heading("20. References")
    add_bullet("Verma, A. (2020). E-Commerce Customer Churn Analysis and Prediction Dataset. Kaggle. https://www.kaggle.com/datasets/ankitverma2010/ecommerce-customer-churn-analysis-and-prediction", "[1] ")
    add_bullet("Fader, P. S., Hardie, B. G., & Lee, K. L. (2005). 'RFM and CLV: Using Iso-value Curves for Customer Base Analysis.' Journal of Marketing Research, 42(4), 415-430.", "[2] ")
    add_bullet("Pedregosa, F., et al. (2011). 'Scikit-learn: Machine Learning in Python.' Journal of Machine Learning Research, 12, 2825-2830.", "[3] ")
    add_bullet("Breiman, L. (2001). 'Random Forests.' Machine Learning, 45(1), 5-32.", "[4] ")
    add_bullet("Reichheld, F. F. (1996). The Loyalty Effect: The Hidden Force Behind Growth, Profits, and Lasting Value. Harvard Business School Press.", "[5] ")
    add_bullet("IBM SkillsBuild & BharatCares (2026). Data Analytics with AI Internship Curriculum & Guidelines, AICTE Association.", "[6] ")

    # Save document
    output_filename = "StudentName_ProjectReport.docx"
    doc.save(output_filename)
    print(f"Successfully generated academic project report: {output_filename}")

if __name__ == '__main__':
    create_report()
