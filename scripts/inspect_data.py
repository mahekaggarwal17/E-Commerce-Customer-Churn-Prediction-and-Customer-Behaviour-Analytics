import openpyxl
import pandas as pd

excel_path = "data/dataset.xlsx"
wb = openpyxl.load_workbook(excel_path, read_only=True)
print("Sheet names:", wb.sheetnames)

sheet = "E Comm" if "E Comm" in wb.sheetnames else wb.sheetnames[0]
print(f"Loading sheet: {sheet}")
df = pd.read_excel(excel_path, sheet_name=sheet)

print("\n--- Basic Info ---")
print(f"Shape: {df.shape}")
print(f"Columns: {list(df.columns)}")

print("\n--- Missing Values ---")
print(df.isnull().sum()[df.isnull().sum() > 0])

print("\n--- Target Variable Distribution ---")
if "Churn" in df.columns:
    print(df["Churn"].value_counts(normalize=True))
    print(df["Churn"].value_counts())

# Save as data/dataset.csv for faster reading and standard CSV format
csv_path = "data/dataset.csv"
df.to_csv(csv_path, index=False)
print(f"\nSaved CSV to: {csv_path}")

print("\n--- Summary Statistics ---")
print(df.describe().T[["count", "mean", "std", "min", "50%", "max"]])
